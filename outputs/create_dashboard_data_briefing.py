from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = r"outputs/dashboard-data-elements-briefing.pptx"
prs = Presentation()
prs.slide_width = Inches(8.27)
prs.slide_height = Inches(11.69)
WHITE=RGBColor(255,255,255); INK=RGBColor(22,28,36); MUTED=RGBColor(85,96,108)
BLUE=RGBColor(45,125,204); BLUE_LIGHT=RGBColor(231,244,252); GRAY=RGBColor(244,246,248); RULE=RGBColor(190,198,207)

def text(slide, value, x, y, w, h, size=15, color=INK, bold=False, align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=MSO_ANCHOR.TOP
    p=tf.paragraphs[0]; p.alignment=align; p.space_after=Pt(0); r=p.add_run(); r.text=value
    r.font.name='Malgun Gothic'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return sh

def rule(slide,x1,y1,x2,y2,color=RULE,width=1.0):
    ln=slide.shapes.add_connector(1,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); ln.line.color.rgb=color; ln.line.width=Pt(width)

def box(slide,x,y,w,h,fill=WHITE,line=WHITE):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.color.rgb=line; sh.line.width=Pt(.8); return sh

def page(title, subtitle):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=WHITE
    text(s,title,.62,.48,7,.5,26,INK,True); text(s,subtitle,.64,1.12,6.9,.3,13,MUTED); rule(s,.64,1.58,7.63,1.58,BLUE,1.5); return s

# 1페이지
s=page('대시보드 데이터 항목 해설 (1/2)','실제 Unity 세션에서 수집되는 값을 기준으로, 한 기록의 연결 관계를 설명합니다.')
text(s,'한 기록의 연결 구조',.64,1.88,7,.3,19,INK,True); box(s,.64,2.32,6.98,.88,BLUE_LIGHT,BLUE)
text(s,'사용자 ID  →  세션  →  모드·시나리오  →  단계  →  PPE  →  결과·시간',.86,2.57,6.55,.25,15,INK,True,PP_ALIGN.CENTER)
text(s,'같은 sessionId를 기준으로 모든 이벤트를 한 사람의 수행 흐름으로 연결합니다.',.86,2.88,6.55,.18,11,MUTED,False,PP_ALIGN.CENTER)
text(s,'01  사용자·세션',.64,3.62,7,.3,19,INK,True)
text(s,'metaAppScopedUserId   Meta 앱 범위에서 사용자를 구분하는 ID\nmetaAgeCategory       Meta가 반환한 연령 범주 (현재 테스트에서는 Unknown 가능)\nsessionId              한 번의 플레이를 묶는 고유 ID\ntimestampUtc           각 이벤트가 발생한 시각\nsession_ended / note   종료 시각과 종료 기록 (예: application_quitting)',.64,4.05,6.95,1.32,14,INK)
text(s,'02  과정·시나리오·단계',.64,5.82,7,.3,19,INK,True)
text(s,'mode                  교육·훈련·테스트 모드\nworkPlan              선택한 시나리오 (예: LeakResponse)\nflowState              현재 진행 단계 또는 상태\n단계별 시작·종료 시각   어느 단계에서 얼마나 머물렀는지 확인하는 근거',.64,6.25,6.95,1.16,14,INK)
text(s,'03  PPE 수행 기록',.64,7.72,7,.3,19,INK,True)
text(s,'itemType              실제 PPE 종류 (왼손·오른손 장갑 등 개별 구분)\ncondition             Clean 또는 Contaminated 상태\n검사 이벤트            PPE 검사 시작·종료 및 결과\nchoice / result        사용 선택과 착용 확인 결과\nrequiredPpeCheck       시나리오 필수 PPE 목록을 모두 착용했는지 여부',.64,8.15,6.95,1.32,14,INK)
box(s,.64,9.85,6.98,.92,GRAY,RULE); text(s,'이 페이지의 핵심',.88,10.08,1.75,.22,14,BLUE,True)
text(s,'사용자 ID와 sessionId가 있어야 한 사람의 시나리오·단계·PPE 수행을 하나로 추적할 수 있습니다.',2.52,10.08,4.72,.38,13,INK)

# 2페이지
s=page('대시보드 데이터 항목 해설 (2/2)','시간·음성·필수 장비 상태와, 현재 데이터로 계산 가능한 지표의 범위를 구분합니다.')
text(s,'04  필수 PPE·음성·시간',.64,1.88,7,.3,19,INK,True)
text(s,'missingRequiredPpe     완료되지 않은 필수 PPE 목록\naudioClip              재생한 교육 음원\naudioLengthSec         음원 전체 길이\naudioElapsedSec        실제 재생된 시간\ncompleted / stopped / replaced  끝까지 재생·중단·다른 음원으로 교체된 상태\n단계·PPE 소요 시간      각 이벤트의 timestampUtc 차이로 계산',.64,2.31,6.95,1.55,14,INK)
text(s,'05  현재 데이터로 확인 가능한 것',.64,4.28,7,.3,19,INK,True)
box(s,.64,4.72,3.34,2.05,GRAY,RULE); text(s,'계산 가능',.88,4.98,2.7,.22,16,BLUE,True)
text(s,'• 시나리오별 완료율\n• 중도 종료율\n• 단계별 소요 시간\n• PPE별 선택·검사 결과\n• 음원 완료·중단 현황\n• 필수 PPE 누락 여부',.88,5.34,2.72,1.2,13,INK)
box(s,4.28,4.72,3.34,2.05,BLUE_LIGHT,BLUE); text(s,'추가 계측이 필요한 것',4.52,4.98,2.9,.22,16,BLUE,True)
text(s,'• 그랩 시도·실패 횟수\n• 실패 후 성공까지 걸린 시간\n• 컨트롤러 좌·우 입력\n• 기존 종료 버튼·EXIT Point 경로 기록\n• 음성 스킵 입력 (그랩과 분리)',4.52,5.34,2.8,1.2,13,INK)
text(s,'06  문제 신호를 상세 확인하는 방식',.64,7.22,7,.3,19,INK,True); box(s,.64,7.66,6.98,1.18,WHITE,BLUE)
text(s,'경고 지표 클릭  →  해당 sessionId 필터  →  시나리오·단계·PPE 이벤트 확인  →  원본 시각·결과 검토',.9,8.02,6.48,.25,14,INK,True,PP_ALIGN.CENTER)
text(s,'정상 기록도 같은 방식으로 상세 조회되어야 비교와 관리가 가능합니다.',.9,8.43,6.48,.18,12,MUTED,False,PP_ALIGN.CENTER)
box(s,.64,9.35,6.98,1.08,GRAY,RULE); text(s,'실제 검증 예시',.88,9.62,1.55,.22,14,BLUE,True)
text(s,'Education · LeakResponse · Meta ID 기록 · 필수 PPE complete 확인\n※ EXIT Point는 중도 중단 수단이며, 종료 버튼과 application_quitting은 별도 종료 경로로 기록해야 합니다.',2.28,9.57,4.95,.48,12,INK)

prs.save(OUT); print(OUT)
