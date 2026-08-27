from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Docs" / "PPT" / "화학물질_안전훈련_VR_PPE착용교육_프로토타입기획서_v0.7_실무형.pptx"

# 기존 세로형 원본과 동일한 페이지 크기(EMU)
SLIDE_W = 7_562_850
SLIDE_H = 10_688_638
W = SLIDE_W / 914_400
H = SLIDE_H / 914_400

FONT = "Malgun Gothic"
NAVY = RGBColor(24, 50, 74)
NAVY_DARK = RGBColor(15, 35, 55)
BLUE = RGBColor(36, 107, 254)
TEAL = RGBColor(25, 145, 120)
AMBER = RGBColor(214, 139, 27)
RED = RGBColor(205, 69, 69)
INK = RGBColor(17, 24, 39)
SLATE = RGBColor(75, 85, 99)
MUTED = RGBColor(107, 114, 128)
BORDER = RGBColor(214, 222, 231)
BG = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(232, 241, 255)
PALE_TEAL = RGBColor(234, 247, 240)
PALE_AMBER = RGBColor(255, 244, 222)
PALE_RED = RGBColor(253, 236, 236)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
prs.core_properties.title = "화학물질 안전훈련 VR · PPE 착용 교육 프로토타입 기획서"
prs.core_properties.subject = "현재 Unity 프로젝트 구현 기준 실무형 프로토타입 기획서"
prs.core_properties.author = "최진영"
prs.core_properties.comments = "원본 PPTX와 분리해 작성한 세로형 실무 기획서"


def set_bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def shape(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=False):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(0.8)
    return shp


def line(slide, x, y, w, color=BORDER, height=0.012):
    return shape(slide, x, y, w, height, color, color)


def textbox(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=12,
    color=INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.03,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=11.5, color=INK, gap=5, symbol="•"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"{symbol}  {item}"
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
    return box


def header(slide, no, title, section="PROTOTYPE PLAN", subtitle=None):
    set_bg(slide)
    shape(slide, 0, 0, W, 0.13, NAVY, NAVY)
    textbox(slide, section, 0.48, 0.31, 4.8, 0.24, 8.5, BLUE, True)
    textbox(slide, title, 0.48, 0.67, 7.25, 0.54, 23, NAVY_DARK, True)
    if subtitle:
        textbox(slide, subtitle, 0.5, 1.21, 7.15, 0.38, 10.5, MUTED)
        y = 1.62
    else:
        y = 1.39
    line(slide, 0.48, y, 7.3, BORDER)
    textbox(slide, f"{no:02d}", 7.22, 0.32, 0.55, 0.25, 9, NAVY, True, PP_ALIGN.RIGHT)
    return y + 0.23


def footer(slide):
    line(slide, 0.48, 11.16, 7.3, BORDER)
    textbox(slide, "화학물질 안전훈련 VR · PPE 착용 교육 프로토타입 기획서", 0.48, 11.23, 5.9, 0.22, 7.5, MUTED)
    textbox(slide, "v0.7", 7.05, 11.23, 0.72, 0.22, 7.5, NAVY, True, PP_ALIGN.RIGHT)


def add_slide(title, section="PROTOTYPE PLAN", subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    y = header(slide, len(prs.slides), title, section, subtitle)
    footer(slide)
    return slide, y


def label(slide, value, x, y, w, color=BLUE, pale=PALE_BLUE):
    shp = shape(slide, x, y, w, 0.34, pale, pale, True)
    textbox(slide, value, x + 0.08, y + 0.065, w - 0.16, 0.18, 8.5, color, True, PP_ALIGN.CENTER)
    return shp


def card(slide, title, body, x, y, w, h, accent=BLUE, pale=PALE_BLUE, body_size=11.2):
    shape(slide, x, y, w, h, WHITE, BORDER, True)
    shape(slide, x, y, 0.07, h, accent, accent, True)
    textbox(slide, title, x + 0.22, y + 0.18, w - 0.42, 0.32, 13, NAVY_DARK, True)
    if isinstance(body, list):
        bullets(slide, body, x + 0.22, y + 0.67, w - 0.44, h - 0.83, body_size, SLATE)
    else:
        textbox(slide, body, x + 0.22, y + 0.66, w - 0.44, h - 0.82, body_size, SLATE)
    return pale


def note(slide, title, body, x, y, w, h, accent=AMBER, pale=PALE_AMBER, size=10.8):
    shape(slide, x, y, w, h, pale, pale, True)
    textbox(slide, title, x + 0.18, y + 0.14, w - 0.36, 0.25, 10.5, accent, True)
    textbox(slide, body, x + 0.18, y + 0.48, w - 0.36, h - 0.60, size, INK)


def status_pill(slide, value, x, y, status="done", w=0.9):
    styles = {
        "done": (TEAL, PALE_TEAL),
        "limit": (AMBER, PALE_AMBER),
        "future": (MUTED, RGBColor(238, 241, 245)),
        "risk": (RED, PALE_RED),
    }
    fg, bg = styles[status]
    label(slide, value, x, y, w, fg, bg)


def add_table(slide, rows, cols, x, y, w, h, col_widths=None, font_size=9.5, header_fill=NAVY):
    table_shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for idx, cw in enumerate(col_widths):
            table.columns[idx].width = Inches(cw)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.045)
            cell.margin_bottom = Inches(0.035)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (WHITE if r % 2 else BG)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in tf.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size if r else font_size + 0.3)
                p.font.bold = r == 0
                p.font.color.rgb = WHITE if r == 0 else INK
                p.space_after = Pt(0)
    return table


def set_cell(table, r, c, value, bold=False, color=None, align=PP_ALIGN.LEFT):
    cell = table.cell(r, c)
    cell.text = value
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for run in p.runs:
            run.font.name = FONT
            run.font.bold = bold or r == 0
            if color:
                run.font.color.rgb = color


def flow_step(slide, n, title, body, x, y, w=2.18, h=1.02, color=BLUE):
    shape(slide, x, y, w, h, WHITE, BORDER, True)
    shape(slide, x + 0.14, y + 0.17, 0.44, 0.44, color, color, True)
    textbox(slide, str(n), x + 0.14, y + 0.27, 0.44, 0.18, 9, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, title, x + 0.7, y + 0.16, w - 0.84, 0.26, 11.2, NAVY_DARK, True)
    textbox(slide, body, x + 0.7, y + 0.48, w - 0.84, h - 0.57, 8.9, SLATE)


# 01. 표지
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)
shape(slide, 0, 0, 0.28, H, NAVY, NAVY)
label(slide, "PPE PROTOTYPE · v0.7", 0.66, 0.82, 2.35, BLUE, PALE_BLUE)
textbox(slide, "화학물질 안전훈련 VR", 0.66, 1.55, 6.85, 0.45, 19, NAVY, True)
textbox(slide, "PPE 착용 교육\n프로토타입 기획서", 0.66, 2.17, 6.65, 1.45, 34, NAVY_DARK, True)
textbox(slide, "현재 Unity 프로젝트 구현 기준 · 실무 검토용", 0.7, 3.95, 6.5, 0.38, 13, BLUE, True)
shape(slide, 0.7, 4.7, 6.85, 3.1, BG, BORDER, True)
textbox(slide, "문서 목적", 0.98, 4.98, 1.25, 0.28, 12, NAVY, True)
bullets(slide, [
    "현재 동작하는 PPE 교육 흐름을 요구사항과 상태 기준으로 정리",
    "입력·UI·VOICE·PPE 판정·거울·퀴즈의 완료 조건 명문화",
    "교육모드 안정화와 후속 훈련·테스트·관리자 개발 기준 제공",
], 0.98, 5.48, 6.1, 1.75, 13, INK, 8)
note(slide, "범위", "현재 프로토타입은 PPE 착용 교육모드 한 사이클이다. 밀폐공간 교육은 별도 후속 시나리오로 관리한다.", 0.7, 8.35, 6.85, 1.15, TEAL, PALE_TEAL, 11.5)
textbox(slide, "작성자  최진영", 0.7, 10.62, 2.6, 0.3, 10, MUTED)
textbox(slide, "A4 PORTRAIT", 6.1, 10.62, 1.45, 0.3, 9, MUTED, True, PP_ALIGN.RIGHT)


# 02. 문서 관리
slide, y = add_slide("문서 관리와 사용 기준", "00 DOCUMENT CONTROL")
t = add_table(slide, 7, 2, 0.55, y + 0.05, 7.18, 4.25, [1.55, 5.63], 10.5)
for c, v in enumerate(["항목", "기준"]): set_cell(t, 0, c, v)
data = [
    ("문서 버전", "v0.7 · 원본 프로토타입 기획서와 동일 버전 유지"),
    ("문서 성격", "발표자료가 아닌 개발·QA·검토용 실무 기획서"),
    ("기준 프로젝트", "Final_VR_Tyche_Pivot/client · Unity 6000.4.8f1"),
    ("기준 씬", "client/Assets/Scenes/3_PPE_Room_HandTest_scale_0.unity"),
    ("실행 경로", "0_App → 1_Title → 2_Intro → 6_LoadingScene_0 → PPE Room"),
    ("관리 원칙", "씬·Prefab·Inspector 작성값을 UI와 Transform의 기준값으로 사용"),
]
for r, pair in enumerate(data, 1):
    set_cell(t, r, 0, pair[0], True, NAVY)
    set_cell(t, r, 1, pair[1])
card(slide, "문서 판정 표현", [
    "구현: 현재 프로젝트에 코드·씬·직렬화 연결이 존재",
    "제한: 구현됐지만 보완 또는 실기기 검증이 필요한 상태",
    "후속: 프로토타입 이후 구현할 범위",
], 0.55, y + 4.65, 7.18, 2.2, TEAL, PALE_TEAL, 11.2)
note(slide, "주의", "정적 코드 확인이나 HTML/PPT 문서 검증을 Quest/OpenXR 정상 동작으로 확대 해석하지 않는다.", 0.55, y + 7.05, 7.18, 0.9, RED, PALE_RED, 10.5)


# 03. 프로젝트 현황
slide, y = add_slide("프로젝트 현황 요약", "01 SCOPE", "현재 구현과 이후 개발을 한눈에 구분한다")
items = [
    ("PPE 교육모드", "전체 흐름 구현·보완 중", "done"),
    ("PPE 훈련모드", "선택 UI·이벤트 존재 / 전용 진행 후속", "limit"),
    ("PPE 테스트모드", "선택 UI·이벤트 존재 / 평가·결과 후속", "limit"),
    ("관리자 기능", "결과 조회·이력 관리 후속", "future"),
    ("밀폐공간 진입 전 교육", "별도 시나리오로 분리", "future"),
]
for i, (name, desc, st) in enumerate(items):
    yy = y + 0.05 + i * 1.15
    shape(slide, 0.55, yy, 7.18, 0.9, WHITE, BORDER, True)
    status_pill(slide, {"done":"구현", "limit":"제한", "future":"후속"}[st], 0.75, yy + 0.25, st, 0.75)
    textbox(slide, name, 1.75, yy + 0.17, 2.1, 0.28, 12.3, NAVY_DARK, True)
    textbox(slide, desc, 3.85, yy + 0.18, 3.55, 0.42, 10.5, SLATE)
note(slide, "프로토타입 완료 정의", "앱 실행부터 PPE 교육 완료·카드 복귀까지 사용자 개입으로 한 사이클을 중단 없이 완주하고, 입력·VOICE·거울·퀴즈 상태가 충돌하지 않는 것.", 0.55, y + 6.05, 7.18, 1.25, BLUE, PALE_BLUE, 11)


# 04. 목표와 사용자
slide, y = add_slide("교육 목표와 사용자 정의", "01 SCOPE")
card(slide, "주 사용자", [
    "화학제품 제조공장 신규·전환배치 작업자",
    "PPE 절차 재교육 대상자",
    "VR 사용 경험이 적은 교육생",
], 0.55, y + 0.05, 3.42, 2.55, BLUE, PALE_BLUE, 11.5)
card(slide, "운영 사용자", [
    "사업장 안전관리자",
    "교육 진행·장비 상태 확인 담당자",
    "후속 결과 조회 기능의 사용 주체",
], 4.25, y + 0.05, 3.48, 2.55, TEAL, PALE_TEAL, 11.5)
textbox(slide, "학습 목표", 0.58, y + 3.05, 1.6, 0.32, 14, NAVY_DARK, True)
goals = [
    ("01", "작업정보 확인", "PPE 선택 전에 태블릿 작업정보를 확인한다."),
    ("02", "보호구 상태 판정", "정상·오염·균열 보호구를 구분해 사용 또는 폐기한다."),
    ("03", "필수 착용 완료", "장비별·좌우별 조건과 테이프 보조 조건을 빠짐없이 완료한다."),
    ("04", "착용 상태 최종 확인", "풀장착 모델과 거울을 통해 최종 상태를 확인한다."),
    ("05", "핵심 내용 복습", "5개 퀴즈에서 오답을 교정하고 전 문항을 완료한다."),
]
for i, (n, title, body) in enumerate(goals):
    yy = y + 3.58 + i * 1.02
    shape(slide, 0.58, yy, 0.55, 0.55, NAVY, NAVY, True)
    textbox(slide, n, 0.58, yy + 0.17, 0.55, 0.2, 9, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, title, 1.35, yy + 0.03, 1.75, 0.27, 11.2, NAVY_DARK, True)
    textbox(slide, body, 3.0, yy + 0.03, 4.55, 0.5, 10.4, SLATE)


# 05. 기술 기준
slide, y = add_slide("실행 환경과 기술 기준", "02 SYSTEM")
t = add_table(slide, 9, 3, 0.55, y + 0.05, 7.18, 5.55, [1.5, 2.45, 3.23], 9.5)
for c, v in enumerate(["구분", "현재 기준", "검증 포인트"]): set_cell(t, 0, c, v)
rows = [
    ("Unity", "6000.4.8f1", "컴파일·씬 참조·직렬화 유지"),
    ("렌더링", "URP 17.4.0", "Quest 양안·투명·거울 성능"),
    ("XR", "OpenXR 1.17.1", "Android/Quest와 PC OpenXR"),
    ("Interaction", "XRI 3.4.1", "Ray·Direct·Teleport 입력 분리"),
    ("Hands", "XR Hands 1.7.3", "손 표시·PPE 장갑 상태"),
    ("Input", "Input System 1.19.0", "Trigger·Grip·Joystick 우선순위"),
    ("타깃", "Meta Quest 2", "72Hz 목표·양안·오디오"),
    ("TTS", "ElevenLabs 박현미", "상태별 클립·중첩 방지"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
card(slide, "필수 기술 원칙", [
    "필수 참조 누락은 런타임 자동 수리 대신 명확한 오류로 중단",
    "UI 표현값은 Scene·Prefab·Inspector가 기준이며 런타임은 상태만 변경",
    "XR 입력 성공은 Game View 마우스가 아니라 실물 장치 이벤트로 검증",
], 0.55, y + 5.85, 7.18, 2.0, RED, PALE_RED, 10.7)


# 06. 장면 경로
slide, y = add_slide("앱 실행·장면 전환 경로", "02 SYSTEM")
route = [
    ("0_App", "앱 공통 초기화"),
    ("1_Title", "로고·타이틀"),
    ("2_Intro", "교육 목적·주의"),
    ("6_LoadingScene_0", "비동기 로딩"),
    ("3_PPE_Room_HandTest_scale_0", "PPE 교육 실행"),
]
for i, (name, body) in enumerate(route):
    yy = y + 0.15 + i * 1.45
    shape(slide, 0.85, yy, 6.55, 0.98, WHITE, BLUE if i == 4 else BORDER, True)
    textbox(slide, f"{i+1:02d}", 1.05, yy + 0.28, 0.5, 0.22, 9.5, BLUE, True)
    textbox(slide, name, 1.75, yy + 0.17, 3.6, 0.29, 11.3, NAVY_DARK, True)
    textbox(slide, body, 1.75, yy + 0.52, 4.95, 0.22, 9.5, SLATE)
    if i < len(route) - 1:
        textbox(slide, "↓", 3.85, yy + 1.02, 0.5, 0.32, 16, BLUE, True, PP_ALIGN.CENTER)
note(slide, "로딩 완료 조건", "0~100% 표시가 역행하지 않고 100% 유지 후 PPE 장면을 한 번만 활성화한다. 전환 뒤 XR Origin·입력·오디오가 정상이어야 한다.", 0.55, y + 7.55, 7.18, 1.1, TEAL, PALE_TEAL, 10.5)


# 07. 상태 흐름
slide, y = add_slide("PPE 교육 상태 흐름", "02 SYSTEM", "VOICE·UI·입력 게이트가 동일 상태를 공유한다")
states = [
    ("Welcome", "첫 안내"), ("NameInput", "한글 이름"), ("Controller", "간단 조작"),
    ("Card", "PPE 카드"), ("Modal", "교육 선택"), ("Mode", "교육모드"),
    ("Move", "PPE 구역"), ("Tablet", "작업정보"), ("PPE", "점검·착용"),
    ("Mirror", "5초 확인"), ("Quiz", "5문제"), ("Return", "카드 복귀"),
]
for i, (title, body) in enumerate(states):
    col = i % 3
    row = i // 3
    x = 0.58 + col * 2.45
    yy = y + 0.05 + row * 1.55
    flow_step(slide, i + 1, title, body, x, yy, 2.18, 1.02, TEAL if i >= 7 else BLUE)
    if col < 2:
        textbox(slide, "→", x + 2.18, yy + 0.32, 0.26, 0.25, 12, MUTED, True, PP_ALIGN.CENTER)
note(slide, "상태 전이 원칙", "각 상태는 화면 표시, VOICE, 허용 입력과 완료 조건을 함께 전환한다. 이전 상태의 UI·AudioClip·코루틴이 다음 상태에 남아서는 안 된다.", 0.55, y + 6.55, 7.18, 1.2, BLUE, PALE_BLUE, 10.5)


# 08. 컨트롤러
slide, y = add_slide("PPE 전용 컨트롤러 조작", "03 INPUT")
controls = [
    ("Ray", "UI 지시", "카드·모달·Action Panel·퀴즈를 가리킨다."),
    ("Trigger", "확정", "UI 선택, 태블릿 확인, 조건부 VOICE 스킵."),
    ("Grip", "직접 Grab", "태블릿·마커 PPE를 잡고 재입력으로 놓는다."),
    ("Joystick", "이동", "텔레포트 Ray를 켜고 유효 위치마커를 선택한다."),
    ("A 버튼", "예비 기능", "키보드 옆 안내 표시는 있으나 현재 기능 연결은 없다."),
]
for i, (key, role, body) in enumerate(controls):
    yy = y + 0.05 + i * 1.48
    shape(slide, 0.55, yy, 7.18, 1.15, WHITE, BORDER, True)
    shape(slide, 0.78, yy + 0.2, 1.15, 0.72, NAVY if i < 4 else SLATE, NAVY if i < 4 else SLATE, True)
    textbox(slide, key, 0.78, yy + 0.42, 1.15, 0.22, 10.5, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, role, 2.18, yy + 0.18, 1.35, 0.25, 11.8, NAVY_DARK, True)
    textbox(slide, body, 3.48, yy + 0.18, 3.9, 0.55, 10.3, SLATE)
note(slide, "입력 분리", "Ray hover는 Trigger 클릭 성공의 증거가 아니다. Quest에서는 PointerDown·Up·Click 또는 XRI Select 이벤트를 각각 확인한다.", 0.55, y + 7.65, 7.18, 1.0, RED, PALE_RED, 10.2)


# 09. 입력 우선순위
slide, y = add_slide("Trigger 우선순위와 이동 게이트", "03 INPUT")
textbox(slide, "한 번의 Trigger 입력은 한 소비자만 처리한다", 0.58, y + 0.05, 7.1, 0.35, 14, NAVY_DARK, True)
priority = [
    ("1", "활성 UI 선택", "카드·모달·버튼·Action Panel·퀴즈"),
    ("2", "잡은 태블릿 확인", "체크·서명·도장 단일 시퀀스"),
    ("3", "VOICE 건너뛰기", "상위 소비자가 없고 음성이 재생 중일 때만"),
]
for i, (n, title, body) in enumerate(priority):
    yy = y + 0.8 + i * 1.5
    shape(slide, 0.85, yy, 6.55, 1.02, WHITE, BLUE if i == 0 else BORDER, True)
    shape(slide, 1.05, yy + 0.22, 0.56, 0.56, BLUE, BLUE, True)
    textbox(slide, n, 1.05, yy + 0.39, 0.56, 0.2, 10, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, title, 1.85, yy + 0.18, 2.0, 0.28, 12.2, NAVY_DARK, True)
    textbox(slide, body, 3.75, yy + 0.18, 3.3, 0.48, 10.2, SLATE)
card(slide, "텔레포트 게이트", [
    "Joystick 기반 이동 인터랙터는 UI Ray와 별도 제어",
    "모달·완료 전환 중 비활성",
    "교육 시작 뒤 유효 위치마커에만 착지",
    "현재 텔레포트 자체 화면 페이드는 없음",
], 0.55, y + 5.45, 7.18, 2.35, TEAL, PALE_TEAL, 10.8)


# 10. VOICE/TTS
slide, y = add_slide("VOICE·TTS 운영 기준", "04 AUDIO")
card(slide, "선정 모델", [
    "서비스: ElevenLabs",
    "보이스: 박현미 모델",
    "톤: 뉴트럴 60% · 아나운서 20% · 온화함 20%",
    "비용: Starter 월 6달러",
], 0.55, y + 0.05, 3.45, 2.8, BLUE, PALE_BLUE, 11.2)
card(slide, "선정 근거", [
    "동종 서비스 대비 자연스러운 딕션과 어투",
    "본 VR의 안전교육 톤앤매너에 적합",
    "기본 크레딧과 저비용 구독으로 반복 제작 가능",
], 4.25, y + 0.05, 3.48, 2.8, TEAL, PALE_TEAL, 11.2)
t = add_table(slide, 6, 3, 0.55, y + 3.2, 7.18, 3.65, [1.35, 2.25, 3.58], 9.4)
for c, v in enumerate(["구분", "발생 시점", "처리 기준"]): set_cell(t, 0, c, v)
rows = [
    ("단계 안내", "상태 진입", "상태별 한 번 재생"),
    ("조건 안내", "선행조건 누락", "해당 조건만 보충"),
    ("오답 안내", "잘못된 판정", "모드별 피드백 수준 적용"),
    ("건너뛰기", "Trigger", "UI·태블릿 미소비 시만"),
    ("완료 안내", "최종 정답", "복귀 전환보다 먼저 재생"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
note(slide, "중복 방지", "새 상태 진입 시 이전 클립·로딩·코루틴을 정리하고 동일 클립이 겹쳐 재생되지 않게 한다.", 0.55, y + 7.15, 7.18, 0.85, RED, PALE_RED, 10.2)


# 11. 전체 교육 흐름
slide, y = add_slide("PPE 착용 교육 전체 흐름", "05 EDUCATION FLOW")
flow = [
    ("진입", "Welcome → 이름 입력 → 간단 컨트롤러 안내"),
    ("선택", "PPE 카드 → 상세 모달 → PPE 착용교육 → 교육모드"),
    ("이동", "텔레포트로 PPE 구역 도착"),
    ("문서", "태블릿 Grab → Trigger 1회 확인"),
    ("점검", "마커 PPE Grab → 사용·폐기 판정 → 착용"),
    ("최종", "풀장착 모델 → 거울 5초 확인"),
    ("복습", "5페이지·3지선다 퀴즈"),
    ("종료", "VOICE → 손·Ray 숨김 → 페이드 → 카드 복귀"),
]
for i, (title, body) in enumerate(flow):
    yy = y + 0.02 + i * 1.05
    shape(slide, 0.7, yy, 7.0, 0.78, WHITE, BORDER, True)
    status_pill(slide, f"{i+1:02d}", 0.9, yy + 0.22, "done", 0.55)
    textbox(slide, title, 1.7, yy + 0.12, 1.05, 0.25, 11.2, NAVY_DARK, True)
    textbox(slide, body, 2.75, yy + 0.12, 4.55, 0.42, 10.2, SLATE)


# 12. 진입 흐름
slide, y = add_slide("진입 단계 명세", "05 EDUCATION FLOW", "Welcome · 이름 입력 · Controller Simp")
t = add_table(slide, 5, 5, 0.55, y + 0.05, 7.18, 4.5, [1.05, 1.35, 1.65, 1.55, 1.58], 8.7)
for c, v in enumerate(["상태", "사용자 입력", "처리", "완료 조건", "예외"]): set_cell(t, 0, c, v)
rows = [
    ("Welcome", "없음", "이름·Enter 안내 VOICE", "안내 종료", "Clip 누락 시 다음 상태 금지"),
    ("NameInput", "한글 키·삭제", "조합 문자열 표시", "Enter 제출", "빈 이름 정책 후속 확정"),
    ("Submit", "Enter", "SubmittedName 전달", "이벤트 1회", "Ray fallthrough 차단"),
    ("Controller", "안내 확인", "Ray·Trigger·Grip·Joystick 설명", "001~005 종료", "A는 시각 표시만"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
card(slide, "Controller Simp 화면 매핑", [
    "001~002: 1_Ray · Card/Panel",
    "003: 2_Marker · 직접 Grab",
    "004~005: 3_Ray_T · 텔레포트",
    "미니 가이드는 Main Camera 자식 World Space Canvas",
], 0.55, y + 4.85, 7.18, 2.25, BLUE, PALE_BLUE, 10.7)
note(slide, "현재 제한", "키보드 옆 A 표시는 긴 Controller Edu 진입 기능이 연결되지 않은 비상호작용 안내 요소다.", 0.55, y + 7.35, 7.18, 0.85, AMBER, PALE_AMBER, 10.2)


# 13. 카드/모달/모드
slide, y = add_slide("카드·모달·모드·이동 명세", "05 EDUCATION FLOW")
steps = [
    ("카드 Trigger", "PPE 상세 모달 표시", "카드만 모달을 연다"),
    ("PPE 착용교육", "교육 종류 선택", "즉시 모달을 닫지 않는다"),
    ("교육모드", "모달 종료·이동 안내", "훈련·테스트는 현재 진행하지 않음"),
    ("Joystick 이동", "PPE 위치마커 텔레포트", "모달 중 차단"),
    ("PPE 구역 도착", "태블릿·PPE 단계 활성화", "도착 이벤트 1회"),
]
for i, (trigger, result, rule) in enumerate(steps):
    yy = y + 0.05 + i * 1.45
    shape(slide, 0.55, yy, 7.18, 1.13, WHITE, BORDER, True)
    textbox(slide, trigger, 0.8, yy + 0.18, 1.55, 0.3, 11.5, NAVY_DARK, True)
    textbox(slide, result, 2.4, yy + 0.18, 2.2, 0.45, 10.2, INK)
    label(slide, rule, 4.78, yy + 0.22, 2.6, TEAL if i != 2 else AMBER, PALE_TEAL if i != 2 else PALE_AMBER)
note(slide, "보존 조건", "카드 입력, 모달 XRI 버튼, 텔레포트 입력을 하나의 fallback으로 합치지 않는다. 각 입력 소비자의 실제 이벤트 경로를 유지한다.", 0.55, y + 7.55, 7.18, 1.0, RED, PALE_RED, 10.2)


# 14. 태블릿
slide, y = add_slide("태블릿 작업정보 확인", "05 EDUCATION FLOW")
card(slide, "사용자 행동", [
    "Grip으로 태블릿을 잡는다.",
    "잡은 상태에서 Trigger를 한 번 누른다.",
    "체크·서명·도장 연출이 끝날 때까지 기다린다.",
], 0.55, y + 0.05, 3.45, 2.75, BLUE, PALE_BLUE, 11.5)
card(slide, "시스템 처리", [
    "필수 항목 체크 → 서명 → 도장 순차 실행",
    "시퀀스 중 추가 Trigger 차단",
    "완료 시 IsDocumentCompleted 설정",
    "태블릿 Grab 중 방호복 Grab 차단",
], 4.25, y + 0.05, 3.48, 2.75, TEAL, PALE_TEAL, 11)
t = add_table(slide, 5, 3, 0.55, y + 3.25, 7.18, 3.2, [1.65, 2.5, 3.03], 9.7)
for c, v in enumerate(["상태", "허용 입력", "판정"]): set_cell(t, 0, c, v)
rows = [
    ("미Grab", "Grip", "태블릿 선택"),
    ("Grab", "Trigger", "단일 확인 시퀀스 시작"),
    ("처리 중", "없음", "재입력 무시"),
    ("완료", "Grip Release", "PPE 단계 선행조건 충족"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
note(slide, "제외된 이전 방식", "B 버튼 페이지 이동, 마지막 페이지 전체 확인, 별도 서명 Trigger는 현재 PPE 태블릿 흐름에서 사용하지 않는다.", 0.55, y + 6.8, 7.18, 1.0, AMBER, PALE_AMBER, 10.4)


# 15. PPE 판정 규칙
slide, y = add_slide("PPE 점검·판정 공통 규칙", "05 EDUCATION FLOW")
rules = [
    ("선택", "아이템 마커가 표시된 PPE를 손으로 직접 Grab", "원거리 UI Ray 선택 금지"),
    ("정상", "Action Panel의 사용 선택", "착용 연출 후 풀장착 자식 활성"),
    ("불량", "Action Panel의 폐기 선택", "동일 검사 오브젝트를 깨끗한 상태로 복원"),
    ("오선택", "정상/불량과 다른 버튼 선택", "완료 처리 금지·모드별 오류 피드백"),
    ("Release", "Grip 재입력 또는 처리 취소", "씬에서 작성한 위치·회전·스케일 복원"),
]
t = add_table(slide, 6, 3, 0.55, y + 0.05, 7.18, 4.75, [1.25, 2.8, 3.13], 9.7)
for c, v in enumerate(["구분", "처리", "불변조건"]): set_cell(t, 0, c, v)
for r, row in enumerate(rules, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
card(slide, "상태 전환", [
    "Idle → Marked → Held → Decision",
    "Decision → Used 또는 Discarded",
    "Discarded → CleanReady → Held → Used",
    "Used 상태는 재선택 불가",
], 0.55, y + 5.15, 7.18, 2.15, TEAL, PALE_TEAL, 10.8)
note(slide, "중요", "폐기 뒤 새 프리팹을 생성하지 않는다. 동일 검사 오브젝트의 결함 시각물과 상태를 정상으로 전환한다.", 0.55, y + 7.55, 7.18, 0.9, RED, PALE_RED, 10.3)


# 16. PPE 완료 매트릭스
slide, y = add_slide("PPE 항목·완료 조건 매트릭스", "05 EDUCATION FLOW")
t = add_table(slide, 11, 5, 0.45, y + 0.02, 7.38, 7.35, [1.28, 1.05, 1.55, 1.35, 2.15], 8.2)
for c, v in enumerate(["항목", "좌우", "불량 처리", "장착 표시", "완료 조건"]): set_cell(t, 0, c, v)
rows = [
    ("방호복", "단일", "폐기→정상", "Suit 루트", "사용 완료"),
    ("호흡보호구", "단일", "균열 폐기", "Mirror Only", "사용 완료"),
    ("안전모", "단일", "상태 판정", "Mirror Only", "사용 완료"),
    ("안전대(등지게)", "단일", "상태 판정", "Suit 자식", "사용 완료"),
    ("장갑 L", "좌", "상태 판정", "왼손/자식", "좌 완료"),
    ("장갑 R", "우", "상태 판정", "오른손/자식", "우 완료"),
    ("장화 L", "좌", "오염 폐기", "Suit 자식", "좌 완료"),
    ("장화 R", "우", "오염 폐기", "Suit 자식", "우 완료"),
    ("테이프", "보조", "해당 없음", "보조 상태", "필수 선행조건"),
    ("거울", "최종", "해당 없음", "평면 반사", "5초 확인"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
note(slide, "최종 PPE 완료", "필수 장비의 사용 완료 + 좌·우 장갑·장화 완료 + 테이프 조건 + 태블릿 완료가 모두 충족되어야 거울 최종 확인을 통과할 수 있다.", 0.45, y + 7.65, 7.38, 1.0, BLUE, PALE_BLUE, 10.2)


# 17. Action Panel
slide, y = add_slide("공용 Action Panel 소유권", "06 INTERACTION")
textbox(slide, "PPE마다 패널을 만들지 않고 하나의 패널을 현재 Grab PPE가 소유한다", 0.55, y + 0.05, 7.18, 0.55, 13.5, NAVY_DARK, True)
ownership = [
    ("Grab", "소유 요청", "처리 중이 아니면 현재 PPE로 교체"),
    ("표시", "사용·폐기 버튼", "PPE 상태와 동일 위치 정책 사용"),
    ("판정", "버튼 Trigger", "현재 소유 PPE에만 결과 전달"),
    ("Release", "소유 해제", "패널 숨김·리스너 정리"),
    ("완료", "Used/Discarded", "중복 이벤트 차단 후 소유 해제"),
]
t = add_table(slide, 6, 3, 0.55, y + 0.8, 7.18, 4.65, [1.3, 2.05, 3.83], 9.5)
for c, v in enumerate(["이벤트", "패널 처리", "검증 기준"]): set_cell(t, 0, c, v)
for r, row in enumerate(ownership, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
card(slide, "경쟁 조건", [
    "양손이 서로 다른 PPE를 동시에 Grab",
    "판정 처리 중 새 PPE Grab",
    "이전 PPE Release 뒤 버튼 리스너 잔류",
    "패널 인스턴스 중복 생성",
], 0.55, y + 5.8, 7.18, 2.15, RED, PALE_RED, 10.8)


# 18. 풀장착
slide, y = add_slide("풀장착 모델 결합 로직", "06 INTERACTION")
textbox(slide, "단일 기준 계층", 0.55, y + 0.05, 1.5, 0.3, 13, NAVY_DARK, True)
hier = [
    ("PPE Body Anchor", "몸 추적 기준"),
    ("PPE_A_SuitWear", "방호복 풀장착 루트"),
    ("장비 자식", "마스크·안전모·안전대·장갑·장화"),
]
for i, (name, body) in enumerate(hier):
    x = 0.65 + i * 2.45
    yy = y + 0.65 + i * 0.45
    shape(slide, x, yy, 2.1, 1.15, WHITE, BLUE if i == 0 else BORDER, True)
    textbox(slide, name, x + 0.15, yy + 0.2, 1.8, 0.3, 11.2, NAVY_DARK, True, PP_ALIGN.CENTER)
    textbox(slide, body, x + 0.15, yy + 0.62, 1.8, 0.24, 9.2, SLATE, False, PP_ALIGN.CENTER)
    if i < 2:
        textbox(slide, "→", x + 2.05, yy + 0.36, 0.38, 0.3, 16, BLUE, True, PP_ALIGN.CENTER)
card(slide, "몸 추적", [
    "HMD의 X/Z 위치 추종",
    "Yaw는 완화해 회전 추종",
    "씬에서 작성한 Y 높이 유지",
    "바닥 접지는 루트 Pivot이 아닌 Renderer Bounds 최저점으로 확인",
], 0.55, y + 3.25, 3.45, 3.1, BLUE, PALE_BLUE, 10.8)
card(slide, "장비 표시", [
    "사용 완료 PPE의 대응 자식만 활성화",
    "마스크·안전모는 본체 시야 방지를 위해 Mirror Only",
    "거울이 별도 모델을 이동시키지 않고 동일 계층 상태를 반영",
    "물리적 몸 슬롯·Snap 판정은 사용하지 않음",
], 4.25, y + 3.25, 3.48, 3.1, TEAL, PALE_TEAL, 10.8)
note(slide, "착용 연출", "사용 확정 후 약 1.8초의 공통 착용 연출을 거쳐 검사 오브젝트 처리와 풀장착 자식 활성화를 같은 판정에서 실행한다.", 0.55, y + 6.75, 7.18, 1.0, AMBER, PALE_AMBER, 10.3)


# 19. 거울
slide, y = add_slide("거울 최종 확인 로직", "06 INTERACTION")
card(slide, "진입 조건", [
    "필수 PPE와 테이프 조건 완료",
    "태블릿 IsDocumentCompleted",
    "거울 약 1.2m 이내",
    "사용자 시선이 거울 방향",
], 0.55, y + 0.05, 3.45, 2.8, BLUE, PALE_BLUE, 11.1)
card(slide, "렌더링", [
    "평면 반사 카메라 사용",
    "RenderTexture 640×1472",
    "마스크·안전모 Mirror Only 레이어",
    "본체·반사 방향·Y·바닥 접지 일치",
], 4.25, y + 0.05, 3.48, 2.8, TEAL, PALE_TEAL, 11.1)
textbox(slide, "확인 시퀀스", 0.58, y + 3.25, 1.6, 0.3, 13, NAVY_DARK, True)
seq = [
    ("조건 진입", "게이지 표시"),
    ("5초 관찰", "진행률 증가"),
    ("완료", "퀴즈 활성화"),
]
for i, (a, b) in enumerate(seq):
    x = 0.65 + i * 2.45
    flow_step(slide, i + 1, a, b, x, y + 3.85, 2.15, 1.05, TEAL)
    if i < 2: textbox(slide, "→", x + 2.14, y + 4.18, 0.28, 0.25, 13, MUTED, True, PP_ALIGN.CENTER)
note(slide, "현재 보완점", "현 구현은 조건 진입 뒤 5초 동안 거리·시선 이탈을 계속 재검사하지 않는다. 타이머 중 조건 이탈 시 초기화하도록 보완한다.", 0.55, y + 5.45, 7.18, 1.05, RED, PALE_RED, 10.3)
card(slide, "Quest 검증", [
    "거울 Off/On 성능 비교",
    "양안 위치·방향 일치",
    "주변 시야 깨짐·검은 화면 확인",
    "RenderTexture 해상도 대안 기록",
], 0.55, y + 6.8, 7.18, 1.65, NAVY, PALE_BLUE, 10.4)


# 20. 퀴즈/완료
slide, y = add_slide("퀴즈와 교육 완료", "05 EDUCATION FLOW")
t = add_table(slide, 6, 4, 0.55, y + 0.05, 7.18, 4.3, [1.25, 1.7, 2.0, 2.23], 9.2)
for c, v in enumerate(["구분", "입력", "처리", "완료 기준"]): set_cell(t, 0, c, v)
rows = [
    ("문제 구성", "Ray", "5페이지·보기 3개", "현재 문제만 선택"),
    ("오답", "Trigger", "경고 후 같은 문제 유지", "재선택 가능"),
    ("정답", "Trigger", "완료 피드백 후 자동 이동", "중복 전환 없음"),
    ("최종 정답", "자동", "완료 흐름 1회 호출", "다섯 문제 정답"),
    ("점수", "없음", "현재 산정·저장하지 않음", "후속 모드에서 추가"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
textbox(slide, "완료 전환 순서", 0.58, y + 4.8, 1.8, 0.3, 13, NAVY_DARK, True)
complete = ["완료 VOICE", "손·Ray 숨김", "화면 페이드", "PPE 카드 위치 복귀", "상태 복원"]
for i, item in enumerate(complete):
    x = 0.52 + i * 1.5
    shape(slide, x, y + 5.38, 1.25, 0.95, WHITE, TEAL if i == 0 else BORDER, True)
    textbox(slide, f"{i+1:02d}", x + 0.1, y + 5.55, 0.3, 0.2, 8, TEAL, True)
    textbox(slide, item, x + 0.12, y + 5.85, 1.0, 0.28, 9.2, NAVY_DARK, True, PP_ALIGN.CENTER)
    if i < 4: textbox(slide, "→", x + 1.24, y + 5.69, 0.25, 0.25, 11, MUTED, True, PP_ALIGN.CENTER)
note(slide, "현재 제한", "별도 점수 결과 화면과 PPE 전체 초기화는 교육모드 현재 완료 범위에 포함하지 않는다.", 0.55, y + 6.75, 7.18, 0.9, AMBER, PALE_AMBER, 10.3)


# 21. 예외 처리
slide, y = add_slide("예외·오류 처리 기준", "07 EXCEPTION")
t = add_table(slide, 9, 4, 0.45, y + 0.02, 7.38, 7.0, [1.55, 2.0, 2.15, 1.68], 8.5)
for c, v in enumerate(["상황", "차단 조건", "사용자 피드백", "복구"]): set_cell(t, 0, c, v)
rows = [
    ("빈 이름 제출", "정책 미확정", "입력 안내", "NameInput 유지"),
    ("모달 중 이동", "Teleport Gate", "이동 Ray 비활성", "모달 종료"),
    ("태블릿 미완료", "PPE 최종 완료 차단", "EDU 016 안내", "태블릿 확인"),
    ("마커 없는 PPE", "Grab 금지", "대상 마커 유지", "올바른 대상 선택"),
    ("정상 PPE 폐기", "판정 불일치", "오답 피드백", "같은 PPE 재선택"),
    ("불량 PPE 사용", "판정 불일치", "오답 피드백", "폐기 선택"),
    ("거울 조건 이탈", "5초 미충족", "게이지 초기화", "조건 재진입"),
    ("필수 참조 누락", "초기화 중단", "경로 포함 오류 1회", "Editor에서 명시 수리"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
note(slide, "금지", "fallback, 자동 선택, 자동 완료, 런타임 오브젝트 생성으로 필수 참조 누락을 숨기지 않는다.", 0.45, y + 7.35, 7.38, 0.9, RED, PALE_RED, 10.3)


# 22. 상태 소유권
slide, y = add_slide("상태·데이터 소유권", "08 DATA")
t = add_table(slide, 10, 4, 0.45, y + 0.02, 7.38, 7.45, [1.45, 2.2, 2.15, 1.58], 8.4)
for c, v in enumerate(["영역", "소유자/기준", "주요 상태", "소비자"]): set_cell(t, 0, c, v)
rows = [
    ("교육 흐름", "PPEVoiceFlowDirector", "FlowState·VOICE", "UI·텔레포트"),
    ("모달", "ScenarioDetailModal", "교육·모드 선택", "Voice Flow"),
    ("태블릿", "PPETabletChecklistController", "IsDocumentCompleted", "PPE·Finale"),
    ("PPE 판정", "PPE 장비 Controller", "Held·Used·Discarded", "Action Panel"),
    ("Action Panel", "단일 패널", "현재 PPE 소유권", "판정 Controller"),
    ("착용 시각", "PPE Body Anchor", "장착 자식 활성", "본체·거울"),
    ("거울", "PPEFinaleController", "거리·시선·5초", "Quiz"),
    ("퀴즈", "PPEQuizController", "페이지·정답", "완료 전환"),
    ("UI 표현", "Scene/Prefab/Inspector", "Transform·색·크기", "Runtime Presenter"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
note(slide, "단일 기준 원칙", "한 상태를 두 컴포넌트가 동시에 완료 처리하지 않는다. 상태 소유자는 판정하고 소비자는 표시·입력 게이트만 갱신한다.", 0.45, y + 7.75, 7.38, 0.9, TEAL, PALE_TEAL, 10.2)


# 23. 기능 요구사항
slide, y = add_slide("핵심 기능 요구사항", "09 REQUIREMENTS")
t = add_table(slide, 13, 4, 0.4, y + 0.02, 7.48, 7.95, [0.75, 2.0, 3.0, 1.73], 7.9)
for c, v in enumerate(["ID", "기능", "인수 기준", "상태"]): set_cell(t, 0, c, v)
rows = [
    ("PF-01", "앱·로딩", "지정 장면 경로 1회 전환", "구현"),
    ("PF-02", "이름 입력", "한글 입력·제출 이벤트 1회", "구현/정책"),
    ("PF-03", "컨트롤러 안내", "Ray·Trigger·Grip·Joystick 설명", "구현"),
    ("PF-04", "카드·모달", "카드 Trigger만 열고 모달 중 이동 차단", "구현"),
    ("PF-05", "태블릿", "Trigger 1회 시퀀스·중복 차단", "구현"),
    ("PF-06", "PPE Grab", "마커 직접 Grab·작성 포즈 복원", "구현/검증"),
    ("PF-07", "PPE 판정", "정상 사용·불량 폐기·오선택 차단", "구현"),
    ("PF-08", "풀장착", "완료 자식만 활성·본체/거울 공유", "구현/검증"),
    ("PF-09", "거울", "1.2m·시선·5초·양안", "보완"),
    ("PF-10", "퀴즈", "5페이지·3보기·오답 재시도", "구현"),
    ("PF-11", "완료", "VOICE·페이드·카드 복귀 1회", "구현"),
    ("PF-12", "훈련·테스트·관리자", "모드 정책·결과 조회", "후속"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row):
        set_cell(t, r, c, v, c in (0, 1), NAVY if c in (0, 1) else None, PP_ALIGN.CENTER if c in (0, 3) else PP_ALIGN.LEFT)


# 24. QA
slide, y = add_slide("QA·인수 기준", "10 QA", "정적 확인과 실기기 완료를 분리한다")
t = add_table(slide, 9, 4, 0.45, y + 0.02, 7.38, 6.6, [1.45, 2.25, 2.15, 1.53], 8.6)
for c, v in enumerate(["검증 영역", "통과 기준", "증거", "환경"]): set_cell(t, 0, c, v)
rows = [
    ("컴파일", "C# 오류 0건", "Unity Console", "Editor"),
    ("장면 경로", "중복 없이 PPE 진입", "실행 로그", "Editor/Quest"),
    ("XR 입력", "각 입력 소비자 1회", "이벤트 로그", "Quest"),
    ("PPE", "마커 Grab·판정·복원", "20회 반복", "Editor/Quest"),
    ("풀장착", "장비·좌우 매핑 일치", "본체/거울 비교", "Quest"),
    ("거울", "양안·Y·바닥·성능", "On/Off 비교", "Quest"),
    ("VOICE", "누락·중첩 0건", "전체 청취", "Editor/Quest"),
    ("완주", "앱→완료→복귀", "3회 연속", "Quest"),
]
for r, row in enumerate(rows, 1):
    for c, v in enumerate(row): set_cell(t, r, c, v, c == 0, NAVY if c == 0 else None)
card(slide, "PPE HandTest 완료 조건", [
    "반복 참조·Collider·Rigidbody 자동 수리 오류 없음",
    "카드·모달·텔레포트 상태 순서 일치",
    "PPE는 마커 직접 Grab만 허용",
    "거울 본체와 반사의 방향·Y·바닥 접지 일치",
], 0.45, y + 6.95, 7.38, 1.75, RED, PALE_RED, 10.1)


# 25. 일정
slide, y = add_slide("프로토타입 이후 개발계획", "11 ROADMAP")
phases = [
    ("08.14–08.31", "PPE 교육 안정화", "VOICE·마커 Grab·태블릿·풀장착·거울·퀴즈·완료 복귀", "교육모드 P0 0건"),
    ("09.01–09.04", "PPE 훈련·테스트", "최소 안내·무힌트·오류·시간·결과 정책 구현", "세 모드 분리"),
    ("09.07–09.11", "관리자·QA", "결과·오류·시간 조회와 회귀·성능 테스트", "주요 결함 해소"),
    ("09.14–09.16", "리허설", "실기기 동선·대본·백업 영상", "시연 시간 준수"),
    ("09.17–09.18", "제출·평가", "APK·기획서·R&D·QA·발표자료", "최종 제출"),
]
for i, (date, title, body, done) in enumerate(phases):
    yy = y + 0.05 + i * 1.55
    shape(slide, 0.55, yy, 7.18, 1.22, WHITE, BORDER, True)
    label(slide, date, 0.75, yy + 0.18, 1.35, BLUE if i < 2 else NAVY, PALE_BLUE)
    textbox(slide, title, 2.3, yy + 0.15, 1.55, 0.28, 11.4, NAVY_DARK, True)
    textbox(slide, body, 3.8, yy + 0.15, 2.8, 0.5, 9.4, SLATE)
    textbox(slide, done, 6.55, yy + 0.18, 0.9, 0.45, 8.7, TEAL, True, PP_ALIGN.CENTER)
note(slide, "별도 후속 시나리오", "밀폐공간 진입 전 교육과 사고 체험은 PPE 프로토타입·세 모드·관리자 기능과 분리된 별도 일정으로 산정한다.", 0.55, y + 8.0, 7.18, 0.9, AMBER, PALE_AMBER, 10.3)


# 26. 산출물과 제한사항
slide, y = add_slide("산출물·알려진 제한·검토 결정", "12 HANDOFF")
card(slide, "산출물", [
    "Unity 프로젝트와 Quest APK",
    "PPE 교육·입력·VOICE·R&D 기획 문서",
    "QA 결과와 Quest 실기기 검증 기록",
    "TTS 원본·최종 AudioClip과 파일 이벤트 맵",
    "시연 대본·백업 영상·에셋 출처",
], 0.55, y + 0.05, 3.45, 3.6, BLUE, PALE_BLUE, 10.7)
card(slide, "현재 알려진 제한", [
    "거울 5초 동안 거리·시선 연속 재검사 필요",
    "빈 이름 허용·길이 정책 미확정",
    "교육 점수 결과·전체 Reset 미구현",
    "훈련·테스트는 선택 이후 진행 로직 후속",
    "Quest 양안·오디오·입력 최종 검증 필요",
], 4.25, y + 0.05, 3.48, 3.6, AMBER, PALE_AMBER, 10.7)
textbox(slide, "검토 시 결정할 항목", 0.58, y + 4.1, 2.1, 0.3, 13, NAVY_DARK, True)
decisions = [
    "빈 이름·최대 길이·금칙문자 정책",
    "거울 조건 이탈 시 게이지 즉시 초기화 여부",
    "교육모드 완료 후 Reset 범위와 돌아가기 화면",
    "훈련·테스트 결과 데이터와 관리자 조회 항목",
    "Quest 목표 성능과 거울 RenderTexture 대안 기준",
]
for i, item in enumerate(decisions):
    yy = y + 4.65 + i * 0.72
    shape(slide, 0.65, yy, 0.34, 0.34, WHITE, BORDER, True)
    textbox(slide, item, 1.2, yy - 0.01, 6.15, 0.42, 10.7, INK)
note(slide, "완료 보고 원칙", "정적 확인, Unity Editor 확인, Quest/OpenXR 확인을 분리해 기록하고 실제로 수행하지 않은 상위 검증을 완료로 표시하지 않는다.", 0.55, y + 8.35, 7.18, 0.9, TEAL, PALE_TEAL, 10.3)


OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
print(f"slides={len(prs.slides)} size={prs.slide_width}x{prs.slide_height}")
